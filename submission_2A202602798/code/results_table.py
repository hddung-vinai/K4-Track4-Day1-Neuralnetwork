# results_table.py — lưu kết quả từng lần chạy ra JSON, rồi điền experiments.xlsx từ mẫu
# templates/experiment_table_template.xlsx (giữ nguyên 4 sheet, tên cột và các cột công thức).
from __future__ import annotations

import json
import math
from pathlib import Path

import numpy as np

COLUMNS = ['exp_id', 'group', 'description', 'loss', 'optimizer', 'lr', 'weight_decay', 'batch', 'epochs',
           'hidden', 'dropout', 'clip_norm', 'precision', 'init', 'seed', 'step0_loss', 'best_val_loss',
           'best_epoch', 'final_train_loss', 'final_val_loss', 'val_acc', 'val_macro_f1', 'time_per_epoch_s',
           'peak_mem_MB', 'diverged', 'eval_acc', 'eval_macro_f1', 'figure_file', 'notes']
FORMULA_COLUMNS = {'step0_gap_vs_lnC', 'gap_val_minus_train', 'delta_val_f1_vs_base', 'beyond_noise'}
OPT_DISPLAY = {'sgd': 'SGD', 'sgd_momentum': 'SGD+momentum', 'adam': 'Adam', 'adamw': 'AdamW'}


def _clean(o):
    # numpy / tuple / NaN -> kiểu JSON thuần (NaN, inf -> None)
    if isinstance(o, dict):
        return {k: _clean(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [_clean(v) for v in o]
    if isinstance(o, np.ndarray):
        return _clean(o.tolist())
    if isinstance(o, (np.integer,)):
        return int(o)
    if isinstance(o, (np.floating, float)):
        o = float(o)
        return o if math.isfinite(o) else None
    if isinstance(o, np.bool_):
        return bool(o)
    return o


def save_result(result: dict, results_dir: str = '../results') -> str:
    # Ghi cfg, history, summary (KHÔNG ghi best_state) ra <results_dir>/<exp_id>.json
    Path(results_dir).mkdir(parents=True, exist_ok=True)
    path = Path(results_dir) / f"{result['cfg']['exp_id']}.json"
    payload = {k: result[k] for k in ('cfg', 'history', 'summary')}
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(_clean(payload), f, indent=1, ensure_ascii=False)
    return str(path)


def _restore_nan(o):
    if isinstance(o, dict):
        return {k: _restore_nan(v) for k, v in o.items()}
    if isinstance(o, list):
        return [math.nan if v is None else _restore_nan(v) for v in o]
    return o


def load_result(path: str) -> dict:
    with open(path, encoding='utf-8') as f:
        r = json.load(f)
    r['history'] = _restore_nan(r['history'])   # None -> NaN để vẽ được
    r['cfg']['hidden'] = tuple(r['cfg']['hidden'])
    return r


def load_results(results_dir: str = '../results') -> list[dict]:
    files = sorted(p for p in Path(results_dir).glob('*.json') if not p.name.startswith('eval_result'))
    return [load_result(str(p)) for p in files]


def _num(v, nd=None):
    if v is None:
        return None
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return round(float(v), nd) if nd is not None else v


def to_row(result: dict, eval_scores: dict | None = None, notes: str = '') -> dict:
    # Một dòng của sheet Experiments: cfg + summary (+ eval_acc/eval_macro_f1 chỉ cho baseline và cấu hình cuối)
    c, s = result['cfg'], result['summary']
    row = dict(
        exp_id=c['exp_id'], group=c['group'], description=c['description'],
        loss=c['loss'].upper(), optimizer=OPT_DISPLAY[c['optimizer']], lr=c['lr'],
        weight_decay=c['weight_decay'], batch=c['batch'], epochs=c['epochs'],
        hidden='-'.join(str(h) for h in c['hidden']), dropout=c['dropout'],
        clip_norm='none' if c['clip_norm'] is None else c['clip_norm'],
        precision=c['precision'], init=c['init'], seed=c['seed'],
        step0_loss=_num(s['step0_loss'], 4), best_val_loss=_num(s['best_val_loss'], 4),
        best_epoch=s['best_epoch'], final_train_loss=_num(s['final_train_loss'], 4),
        final_val_loss=_num(s['final_val_loss'], 4), val_acc=_num(s['val_acc'], 4),
        val_macro_f1=_num(s['val_macro_f1'], 4), time_per_epoch_s=_num(s['time_per_epoch_s'], 3),
        peak_mem_MB=_num(s['peak_mem_MB'], 1), diverged='Y' if s['diverged'] else 'N',
        eval_acc=None, eval_macro_f1=None, figure_file=f"figures/{c['exp_id']}.png",
    )
    if eval_scores is not None:
        row['eval_acc'] = round(eval_scores['accuracy'], 4)
        row['eval_macro_f1'] = round(eval_scores['macro_f1'], 4)
    extra = []
    if c.get('scheduler'):
        extra.append(f"scheduler={c['scheduler']}")
    if c.get('notes'):
        extra.append(c['notes'])
    if s['peak_mem_MB'] is None:
        extra.append('peak_mem: không có GPU')
    if notes:
        extra.append(notes)
    row['notes'] = ' | '.join(extra)
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seed_ids: list[str] | None = None, summary_notes: dict | None = None) -> None:
    # Điền rows vào sheet Experiments từ dòng 2 (bỏ qua cột công thức), exp_id baseline vào sheet Seeds,
    # nhận xét vào cột H của sheet Summary. KHÔNG dùng data_only=True (sẽ mất công thức).
    import openpyxl
    wb = openpyxl.load_workbook(template_path)
    ws = wb['Experiments']
    header = {ws.cell(1, j).value: j for j in range(1, ws.max_column + 1)}
    assert all(c in header for c in COLUMNS), 'mẫu thiếu cột'
    assert len(rows) <= 60, 'mẫu có công thức cho 60 dòng'
    for i, row in enumerate(rows, start=2):
        for col in COLUMNS:
            ws.cell(i, header[col]).value = row.get(col)
    if seed_ids is not None:
        ss = wb['Seeds']
        for k in range(5):
            ss.cell(2 + k, 1).value = seed_ids[k] if k < len(seed_ids) else None
    if summary_notes:
        sm = wb['Summary']
        for r in range(2, 12):
            g = sm.cell(r, 1).value
            if g in summary_notes:
                sm.cell(r, 8).value = summary_notes[g]
    wb.save(out_path)
